#!/usr/bin/env python3
"""Generate an editable Word CV (_private/Wilhan_Nunes_CV.docx) from content/cv.toml.

The document is built with python-docx from the same data as the website
(`sitegen.load_cv`). It uses real heading paragraphs, tab-aligned dates and
clickable hyperlinks so it stays ATS-friendly. `_private/` is ignored by Jekyll,
so the file is kept in the repo but never served on the site.
"""

import datetime as dt
import html
import io
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from sitegen import load_cv

DOCX_OUTPUT = Path("_private/Wilhan_Nunes_CV.docx")

FONT = "Calibri"
BODY_SIZE = Pt(10.5)
ACCENT = RGBColor(0x1E, 0x5F, 0xA8)
ACCENT_DARK = RGBColor(0x0F, 0x3D, 0x72)
MUTED = RGBColor(0x6B, 0x6B, 0x72)
MARGIN = Mm(18)
PAGE_WIDTH, PAGE_HEIGHT = Mm(210), Mm(297)
TEXT_WIDTH = PAGE_WIDTH - 2 * MARGIN

TAG_RE = re.compile(r"(</?(?:strong|b|em|i)>)", re.IGNORECASE)
OTHER_TAG_RE = re.compile(r"<[^>]+>")


# ── Low-level helpers ────────────────────────────────────────────────────────

def add_rich_text(paragraph, text: str, *, bold=False, italic=False, color=None, size=None):
    """Append text with inline <strong>/<em> tags as bold/italic runs."""
    state = {"b": bold, "i": italic}
    for chunk in TAG_RE.split(text):
        tag = chunk.lower()
        if tag in ("<strong>", "<b>"):
            state["b"] = True
        elif tag in ("</strong>", "</b>"):
            state["b"] = bold
        elif tag in ("<em>", "<i>"):
            state["i"] = True
        elif tag in ("</em>", "</i>"):
            state["i"] = italic
        elif chunk:
            clean = html.unescape(OTHER_TAG_RE.sub("", chunk))
            if clean:
                run = paragraph.add_run(clean)
                run.bold, run.italic = state["b"], state["i"]
                if color is not None:
                    run.font.color.rgb = color
                if size is not None:
                    run.font.size = size
    return paragraph


def add_hyperlink(paragraph, url: str, text: str, *, color=ACCENT, underline=True):
    """Insert a clickable external hyperlink run into a paragraph."""
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "Hyperlink")
    rpr.append(style)
    col = OxmlElement("w:color")
    col.set(qn("w:val"), str(color))
    rpr.append(col)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def add_field(paragraph, instr: str, *, color=MUTED, size=Pt(8.5)):
    """Insert a simple Word field (e.g. PAGE) as a run sequence."""
    def fld_char(kind):
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        return el

    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = f" {instr} "
    for el in (fld_char("begin"), instr_el, fld_char("separate")):
        run = paragraph.add_run()
        run._r.append(el)
    placeholder = paragraph.add_run("1")
    end = paragraph.add_run()
    end._r.append(fld_char("end"))
    for run in (placeholder,):
        run.font.color.rgb = color
        run.font.size = size


def set_bottom_border(paragraph, color="1E5FA8", size=6):
    """Add a thin rule under a paragraph (used for section headings)."""
    ppr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    ppr.append(borders)


def spacing(paragraph, before=0, after=0, line=None):
    fmt = paragraph.paragraph_format
    fmt.space_before, fmt.space_after = Pt(before), Pt(after)
    if line is not None:
        fmt.line_spacing = line
    return paragraph


def remove_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tbl_pr.append(borders)


# ── Document setup ───────────────────────────────────────────────────────────

def setup_document(cv: dict) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = PAGE_WIDTH, PAGE_HEIGHT
    section.left_margin = section.right_margin = MARGIN
    section.top_margin = section.bottom_margin = MARGIN
    section.footer_distance = Mm(9)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = BODY_SIZE
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, color in (("Title", 22, ACCENT_DARK), ("Heading 1", 12.5, ACCENT_DARK),
                              ("Heading 2", 10.5, ACCENT)):
        style = doc.styles[name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.italic = False
        style.font.color.rgb = color
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rfonts.set(qn(attr), FONT)
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rfonts.attrib.pop(qn(attr), None)
    title_ppr = doc.styles["Title"].element.get_or_add_pPr()
    for bdr in title_ppr.findall(qn("w:pBdr")):  # default Title has a bottom rule
        title_ppr.remove(bdr)

    if "Hyperlink" not in [s.name for s in doc.styles]:
        doc.styles.add_style("Hyperlink", WD_STYLE_TYPE.CHARACTER)

    props = doc.core_properties
    props.author = cv["person"]["full_name"]
    props.title = f"{cv['person']['full_name']} — Curriculum Vitae"
    props.subject = cv["person"]["field"]
    props.keywords = "CV, curriculum vitae, metabolomics"
    props.last_modified_by = cv["person"]["full_name"]
    props.comments = "Generated by build_docx.py from content/cv.toml"
    # Fixed dates so unchanged content gives a byte-identical file (no CI churn).
    props.created = props.modified = dt.datetime(2026, 1, 1)
    return doc


def add_footer(doc: Document, cv: dict) -> None:
    para = doc.sections[0].footer.paragraphs[0]
    para.alignment = 1  # center
    run = para.add_run(f"{cv['person']['full_name']} — CV · page ")
    run.font.color.rgb, run.font.size = MUTED, Pt(8.5)
    add_field(para, "PAGE")


# ── Sections ─────────────────────────────────────────────────────────────────

def section_heading(doc: Document, text: str) -> None:
    para = doc.add_paragraph(text, style="Heading 1")
    spacing(para, before=11, after=4)
    para.paragraph_format.keep_with_next = True
    set_bottom_border(para)


def entry_title(doc: Document, title: str, dates: str):
    """Bold title on the left, muted dates right-aligned via a tab stop."""
    para = doc.add_paragraph()
    para.paragraph_format.tab_stops.add_tab_stop(TEXT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
    para.paragraph_format.keep_with_next = True
    spacing(para, before=5)
    add_rich_text(para, title, bold=True, color=ACCENT_DARK)
    run = para.add_run(f"\t{dates}")
    run.font.color.rgb, run.font.size = MUTED, Pt(9.5)
    return para


def muted_line(doc: Document, text: str, *, italic=False, keep=True):
    para = doc.add_paragraph()
    para.paragraph_format.keep_with_next = keep
    add_rich_text(para, text, italic=italic, color=MUTED, size=Pt(9.5))
    return para


def add_header(doc: Document, cv: dict) -> None:
    p = cv["person"]
    name = doc.add_paragraph(p["full_name"], style="Title")
    spacing(name, after=2, line=1.0)

    tagline = doc.add_paragraph()
    add_rich_text(tagline, " · ".join([p["degree"], p["role"], p["field"]]),
                  bold=True, color=ACCENT, size=Pt(11.5))
    spacing(tagline, after=2)

    contact = doc.add_paragraph()
    add_hyperlink(contact, f"mailto:{p['email']}", p["email"], underline=False)
    if p.get("institutional_email"):
        add_rich_text(contact, "  ·  ", color=MUTED)
        add_hyperlink(contact, f"mailto:{p['institutional_email']}", p["institutional_email"], underline=False)
    add_rich_text(contact, f"  ·  {p['institution']}  ·  {p['location']}", color=MUTED)
    spacing(contact, after=1)

    links = doc.add_paragraph()
    for i, prof in enumerate(cv["profiles"]):
        if i:
            add_rich_text(links, "  ·  ", color=MUTED)
        add_rich_text(links, f"{prof['name']}: ", color=MUTED, size=Pt(9.5))
        add_hyperlink(links, prof["url"], display_url(prof["url"]))
    for run in links.runs:
        run.font.size = Pt(9.5)


def display_url(url: str) -> str:
    """Shorten a URL for display (drop scheme, trailing slash and tracking params)."""
    short = re.sub(r"^https?://(www\.)?", "", url).rstrip("/")
    return short.replace("&hl=en", "")


def add_profile(doc: Document, cv: dict) -> None:
    section_heading(doc, "Profile")
    para = add_rich_text(doc.add_paragraph(), cv["summary"]["text"])
    para.paragraph_format.alignment = 3  # justify


def add_entries(doc: Document, heading: str, entries: list[dict]) -> None:
    section_heading(doc, heading)
    for e in entries:
        entry_title(doc, e["title"], e["dates"])
        has_body = bool(e.get("text") or e.get("bullets"))
        muted_line(doc, e["org"], keep=has_body)
        if e.get("text"):
            add_rich_text(spacing(doc.add_paragraph(), before=1), e["text"])
        for bullet in e.get("bullets", []):
            para = doc.add_paragraph(style="List Bullet")
            spacing(para, before=1)
            add_rich_text(para, bullet)


def add_education(doc: Document, cv: dict) -> None:
    section_heading(doc, "Education")
    for e in cv["education"]:
        entry_title(doc, e["degree"], e["dates"])
        muted_line(doc, e["institution"])
        muted_line(doc, e["advisor"], italic=True, keep=False)


def add_skills(doc: Document, cv: dict) -> None:
    """Borderless two-column table: group label | comma-separated items."""
    section_heading(doc, "Skills")
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    remove_table_borders(table)
    widths = (Mm(30), TEXT_WIDTH - Mm(30))
    for skill in cv["skills"]:
        cells = table.add_row().cells
        for cell, width in zip(cells, widths):
            cell.width = width
        label, items = cells[0].paragraphs[0], cells[1].paragraphs[0]
        add_rich_text(label, skill["group"], bold=True, color=ACCENT_DARK)
        add_rich_text(items, ", ".join(skill["items"]))
        for para in (label, items):
            spacing(para, before=2, after=2)


def format_authors(paragraph, authors: list[dict]) -> None:
    parts = []
    for a in authors:
        if a.get("ellipsis"):
            parts.append(("…", False))
        elif a.get("etal"):
            parts.append(("et al.", False))
        else:
            parts.append((a["name"], a["me"]))
    for i, (text, me) in enumerate(parts):
        if i:
            paragraph.add_run(", ")
        run = paragraph.add_run(text)
        run.bold = me


def add_publication(doc: Document, number: int, pub: dict) -> None:
    para = doc.add_paragraph()
    fmt = para.paragraph_format
    fmt.left_indent, fmt.first_line_indent = Mm(8), Mm(-8)
    fmt.tab_stops.add_tab_stop(Mm(8))
    fmt.keep_together = True
    spacing(para, before=1, after=4)

    num = para.add_run(f"{number}.\t")
    num.font.color.rgb = MUTED
    format_authors(para, pub["authors"])
    # Initials and "et al." already end in a period; avoid a doubled "..".
    para.add_run(" " if para.text.endswith(".") else ". ")
    add_rich_text(para, f"{pub['title']}. ")
    if pub["journal"]:
        add_rich_text(para, pub["journal"], italic=True)
    details = [x for x in (pub["volume"], pub["pages"]) if x]
    tail = (", " + ", ".join(details)) if details and pub["journal"] else ", ".join(details)
    para.add_run(f"{tail} ({pub['year']}).")
    if pub["doi"]:
        para.add_run(" ")
        add_hyperlink(para, f"https://doi.org/{pub['doi']}", f"doi:{pub['doi']}")


def add_publications(doc: Document, cv: dict) -> None:
    counts = cv["pub_counts"]
    section_heading(doc, "Publications")
    summary = doc.add_paragraph()
    add_rich_text(summary, f"{counts['article']} peer-reviewed articles · {counts['preprint']} preprints · "
                           f"{counts['first_author']} as first author. Own name in bold.",
                  color=MUTED, size=Pt(9.5))
    spacing(summary, after=2)

    for kind, label in (("article", "Peer-reviewed articles"), ("preprint", "Preprints")):
        pubs = [p for p in cv["publications"] if p["type"] == kind]
        if not pubs:
            continue
        sub = doc.add_paragraph(label, style="Heading 2")
        spacing(sub, before=6, after=3)
        sub.paragraph_format.keep_with_next = True
        # Reverse-chronological list, numbered down so the oldest item is 1.
        for i, pub in enumerate(pubs):
            add_publication(doc, len(pubs) - i, pub)


def save_reproducible(doc: Document, path: Path) -> None:
    """Save with fixed zip timestamps so identical content gives identical bytes."""
    buf = io.BytesIO()
    doc.save(buf)
    with zipfile.ZipFile(buf) as src, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            dst.writestr(zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0)),
                         src.read(info.filename), compress_type=zipfile.ZIP_DEFLATED)


# ── Main ─────────────────────────────────────────────────────────────────────

def build_docx() -> None:
    cv = load_cv()
    doc = setup_document(cv)
    add_footer(doc, cv)
    add_header(doc, cv)
    add_profile(doc, cv)
    add_entries(doc, "Research", cv["research"])
    add_entries(doc, "Professional Experience", cv["experience"])
    add_education(doc, cv)
    add_skills(doc, cv)
    add_publications(doc, cv)

    DOCX_OUTPUT.parent.mkdir(exist_ok=True)
    save_reproducible(doc, DOCX_OUTPUT)
    print(f"✓ Generated {DOCX_OUTPUT}")


if __name__ == "__main__":
    build_docx()
